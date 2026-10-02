#!/usr/bin/env python3
"""
kr-redev-now 정적 스냅샷 수집기.

입력:
  - scripts/raw/files/OA-22712_도시계획사업_서울플랜_202609.zip (구역 폴리곤 shp + 코드표 xlsx)
  - 서울 열린데이터광장 OpenAPI: TbSeoulRedevStatus, ListCombinationBidding, CleanupPartnersInfo

출력 (data/):
  - zones.geojson, projects.json, bids.json, bizzones.json, partners.json, meta.json

런타임에는 API를 호출하지 않는다. 이 스크립트만 네트워크를 쓴다.
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import time
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
import shapefile  # pyshp
from pyproj import CRS, Transformer
from shapely.geometry import mapping, shape
from shapely.ops import transform as shapely_transform

ROOT = Path(__file__).resolve().parent.parent
RAW_FILES = ROOT / "scripts" / "raw" / "files"
DATA_DIR = ROOT / "data"
ZIP_PATH = RAW_FILES / "OA-22712_도시계획사업_서울플랜_202609.zip"
ENV_PATH = ROOT / ".env.local"

TIMEOUT_SEC = 15
RETRIES = 3
PAGE_SIZE = 1000

# ---------------------------------------------------------------------------
# 환경변수 / API 키
# ---------------------------------------------------------------------------


def load_seoul_api_key() -> str:
    if not ENV_PATH.exists():
        sys.exit(f"환경파일 없음: {ENV_PATH}")
    for line in ENV_PATH.read_text().splitlines():
        line = line.strip()
        if line.startswith("SEOUL_API_KEY="):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    sys.exit("SEOUL_API_KEY 를 .env.local 에서 찾지 못함")


# ---------------------------------------------------------------------------
# 서울 열린데이터광장 OpenAPI 호출 (타임아웃 15초 + 재시도 3회)
# ---------------------------------------------------------------------------


def fetch_page(api_key: str, service: str, start: int, end: int) -> dict:
    url = f"http://openapi.seoul.go.kr:8088/{api_key}/json/{service}/{start}/{end}/"
    last_err = None
    for attempt in range(1, RETRIES + 1):
        try:
            resp = requests.get(url, timeout=TIMEOUT_SEC)
            resp.raise_for_status()
            data = resp.json()
            body = data.get(service)
            if body is None:
                raise RuntimeError(f"응답에 {service} 키 없음: {list(data.keys())}")
            code = body.get("RESULT", {}).get("CODE")
            if code != "INFO-000":
                raise RuntimeError(f"RESULT.CODE={code}: {body.get('RESULT')}")
            return body
        except Exception as e:  # noqa: BLE001 - 네트워크 장애 전부 재시도 대상
            last_err = e
            print(f"  [재시도 {attempt}/{RETRIES}] {service} {start}-{end}: {e}", file=sys.stderr)
            time.sleep(1.5 * attempt)
    raise RuntimeError(f"{service} {start}-{end} 호출 실패: {last_err}")


def fetch_all(api_key: str, service: str, total_hint: int | None = None) -> list[dict]:
    rows: list[dict] = []
    start = 1
    total = total_hint or PAGE_SIZE
    while True:
        end = start + PAGE_SIZE - 1
        body = fetch_page(api_key, service, start, end)
        total = int(body.get("list_total_count", total))
        page_rows = body.get("row") or []
        rows.extend(page_rows)
        print(f"  {service}: {len(rows)}/{total}")
        if len(page_rows) < PAGE_SIZE or len(rows) >= total:
            break
        start += PAGE_SIZE
    return rows


# ---------------------------------------------------------------------------
# 코드표 (xlsx) 파싱
# ---------------------------------------------------------------------------


def extract_zip_members(zip_path: Path) -> dict[str, bytes]:
    """zip 내부 파일명이 cp437로 깨져 들어오므로 cp949 로 복원해 basename 기준 dict 반환."""
    out: dict[str, bytes] = {}
    with zipfile.ZipFile(zip_path) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            try:
                name = name.encode("cp437").decode("cp949")
            except Exception:
                pass
            out[os.path.basename(name)] = z.read(info)
    return out


def parse_code_tables(xlsx_bytes: bytes) -> tuple[dict, dict, dict, set]:
    """
    반환:
      lclas_label: LCLAS_CL -> 대분류 한글명
      sclas_label: SCLAS_CL -> 소분류 한글명
      propel_label: PROPEL_CD -> (사업구분, 중분류, 추진단계 한글명)
      all_propel_codes: 코드표에 정의된 전체 PROPEL_CD 집합
    """
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    ws1 = wb["01. 사업유형_코드"]
    lclas_label: dict[str, str] = {}
    sclas_label: dict[str, str] = {}
    cur_lclas_label = None
    for row in ws1.iter_rows(min_row=2, values_only=True):
        lclas_name, lclas_cl, _mid_name, sclas_name, sclas_cl = row
        if lclas_name:
            cur_lclas_label = lclas_name.strip()
        if lclas_cl and cur_lclas_label:
            lclas_label[lclas_cl.strip()] = cur_lclas_label
        if sclas_cl and sclas_name:
            sclas_label[sclas_cl.strip()] = sclas_name.strip()

    ws2 = wb["02. 추진단계_코드"]
    propel_label: dict[str, tuple[str, str, str]] = {}
    cur_biz_gubun = None
    cur_midcat = None
    for row in ws2.iter_rows(min_row=2, values_only=True):
        biz_gubun, midcat, propel_cd, label = row[0], row[1], row[2], row[3]
        if biz_gubun:
            cur_biz_gubun = biz_gubun.strip()
        if midcat:
            cur_midcat = midcat.strip()
        if propel_cd and label:
            propel_label[propel_cd.strip()] = (cur_biz_gubun, cur_midcat, label.strip())

    return lclas_label, sclas_label, propel_label, set(propel_label.keys())


# ---------------------------------------------------------------------------
# 자치구 코드 매핑 (법정동 시군구 코드, 11110~11740)
# ---------------------------------------------------------------------------

GU_CODE_MAP = {
    "11110": "종로구",
    "11140": "중구",
    "11170": "용산구",
    "11200": "성동구",
    "11215": "광진구",
    "11230": "동대문구",
    "11260": "중랑구",
    "11290": "성북구",
    "11305": "강북구",
    "11320": "도봉구",
    "11350": "노원구",
    "11380": "은평구",
    "11410": "서대문구",
    "11440": "마포구",
    "11470": "양천구",
    "11500": "강서구",
    "11530": "구로구",
    "11545": "금천구",
    "11560": "영등포구",
    "11590": "동작구",
    "11620": "관악구",
    "11650": "서초구",
    "11680": "강남구",
    "11710": "송파구",
    "11740": "강동구",
}


# ---------------------------------------------------------------------------
# 구역명 정규화 (조인 키)
# ---------------------------------------------------------------------------

_BRACKET_RE = re.compile(r"[\(\[\{（][^)\]\}）]*[\)\]\}）]")
_WS_RE = re.compile(r"\s+")
_JE_BEFORE_DIGIT_RE = re.compile(r"제(?=\d)")

# 긴 복합어부터 제거해야 부분 잔재가 안 남는다.
_SUFFIX_TOKENS = [
    "주택재개발정비사업조합",
    "주택재건축정비사업조합",
    "도시정비형재개발정비사업조합",
    "도시환경정비사업조합",
    "가로주택정비사업조합",
    "소규모재건축사업조합",
    "소규모재개발사업조합",
    "조합설립추진위원회",
    "설립추진위원회",
    "준비위원회",
    "정비사업조합",
    "재개발조합",
    "재건축조합",
    "추진위원회",
    "주택재개발",
    "주택재건축",
    "도시정비형",
    "도시환경",
    "정비사업",
    "재개발",
    "재건축",
    "추진위",
    "조합설립",
    "조합",
    "주택",
    "아파트",
    "일대",
    "번지",
    "구역",
    "사업",
]


def normalize_zone_name(name: str | None) -> str:
    if not name:
        return ""
    s = name.strip()
    s = _BRACKET_RE.sub("", s)
    s = _WS_RE.sub("", s)
    s = _JE_BEFORE_DIGIT_RE.sub("", s)
    for tok in _SUFFIX_TOKENS:
        s = s.replace(tok, "")
    s = s.replace("-", "")
    return s.lower()


# ---------------------------------------------------------------------------
# 폴리곤 (shp) 파싱 + 좌표 변환 + 단순화
# ---------------------------------------------------------------------------


def round_geojson_coords(obj, ndigits=6):
    if isinstance(obj, (list, tuple)):
        if obj and isinstance(obj[0], (int, float)):
            return [round(v, ndigits) for v in obj]
        return [round_geojson_coords(x, ndigits) for x in obj]
    return obj


def build_zones(shp_bytes, shx_bytes, dbf_bytes, prj_text, lclas_label, sclas_label, propel_label):
    crs = CRS.from_wkt(prj_text)
    transformer = Transformer.from_crs(crs, CRS.from_epsg(4326), always_xy=True)

    def to_wgs84(x, y, z=None):
        lon, lat = transformer.transform(x, y)
        return (lon, lat)

    sf = shapefile.Reader(
        shp=io.BytesIO(shp_bytes),
        shx=io.BytesIO(shx_bytes),
        dbf=io.BytesIO(dbf_bytes),
        encoding="cp949",
    )
    field_names = [f[0] for f in sf.fields[1:]]

    features = []
    missing_propel_codes = Counter()
    simplify_before_bytes = 0
    simplify_after_bytes = 0
    topology_breaks = 0

    for sr in sf.iterShapeRecords():
        rec = dict(zip(field_names, sr.record))
        geom = sr.shape.__geo_interface__
        try:
            shp_geom = shape(geom)
        except Exception as e:
            print(f"  [경고] geometry 파싱 실패 PRESENT_SN={rec.get('PRESENT_SN')}: {e}", file=sys.stderr)
            continue
        if not shp_geom.is_valid:
            shp_geom = shp_geom.buffer(0)

        wgs_geom = shapely_transform(to_wgs84, shp_geom)

        before_json = json.dumps(mapping(wgs_geom))
        simplify_before_bytes += len(before_json.encode("utf-8"))

        # 단순화: 약 1m 허용오차를 degree 로 환산 (위도 37.5 부근 1m ~ 0.000009도)
        tol_deg = 1.0 / 111_320.0
        simplified = wgs_geom.simplify(tol_deg, preserve_topology=True)
        if simplified.is_empty or not simplified.is_valid:
            simplified = wgs_geom
            topology_breaks += 1

        after_json_obj = mapping(simplified)
        after_json_obj = {
            **after_json_obj,
            "coordinates": round_geojson_coords(after_json_obj["coordinates"]),
        }
        after_bytes = json.dumps(after_json_obj).encode("utf-8")
        simplify_after_bytes += len(after_bytes)

        try:
            rep_point = simplified.representative_point()
        except Exception:
            rep_point = wgs_geom.representative_point()
        lng, lat = round(rep_point.x, 6), round(rep_point.y, 6)

        propel_cd = (rec.get("PROPEL_CD") or "").strip()
        stage_label = None
        if propel_cd:
            entry = propel_label.get(propel_cd)
            if entry:
                stage_label = entry[2]
            else:
                missing_propel_codes[propel_cd] += 1

        sigungu = (rec.get("SIGNGU_SE") or "").strip()
        gu_name = GU_CODE_MAP.get(sigungu)

        lclas_cl = (rec.get("LCLAS_CL") or "").strip()
        sclas_cl = (rec.get("SCLAS_CL") or "").strip()

        props = {
            "id": rec.get("PRESENT_SN"),
            # 문자열 필드는 결측이어도 "" — 앱이 문자열 메서드를 바로 쓴다(코드표에 없는 PP1902 등)
            "name": rec.get("DGM_NM") or "",
            "gu": gu_name or "",
            "guCode": sigungu or None,
            "cat": (lclas_label.get(lclas_cl) or "").strip(),
            "type": (sclas_label.get(sclas_cl) or "").strip(),
            "stage": (stage_label or "").strip(),
            "stageCode": propel_cd or None,
            "area": rec.get("DGM_AR"),
            "ntfc": rec.get("NTFC_SN") or None,
            "grp": rec.get("GRP") or None,
            "lat": lat,
            "lng": lng,
            "statId": None,  # 조인 1에서 채움
            "bizNos": [],  # 조인 2에서 채움
        }

        features.append(
            {
                "type": "Feature",
                "geometry": after_json_obj,
                "properties": props,
            }
        )

    stats = {
        "missing_propel_codes": dict(missing_propel_codes),
        "simplify_before_bytes": simplify_before_bytes,
        "simplify_after_bytes": simplify_after_bytes,
        "topology_breaks": topology_breaks,
    }
    return features, stats


# ---------------------------------------------------------------------------
# 조인 1: 통계(TbSeoulRedevStatus) <-> 폴리곤(zones)
# ---------------------------------------------------------------------------


def join_stats_to_zones(stat_rows, zone_features):
    zones_by_gu: dict[str, list[dict]] = defaultdict(list)
    for feat in zone_features:
        gu = feat["properties"]["gu"]
        if gu:
            zones_by_gu[gu].append(feat)

    matched = 0
    unmatched_samples = []
    ambiguous_samples = []
    stat_zone_id: dict[str, str | None] = {}

    for row in stat_rows:
        seq_no = row.get("SEQ_NO")
        gu = (row.get("DISTRICT") or "").strip()
        zone_nm = row.get("ZONE_NM") or ""
        norm_target = normalize_zone_name(zone_nm)
        candidates = zones_by_gu.get(gu, [])

        exact = [f for f in candidates if normalize_zone_name(f["properties"]["name"]) == norm_target and norm_target]
        contain = []
        if not exact and norm_target:
            for f in candidates:
                norm_cand = normalize_zone_name(f["properties"]["name"])
                if not norm_cand:
                    continue
                if norm_target in norm_cand or norm_cand in norm_target:
                    contain.append(f)

        chosen = None
        if len(exact) == 1:
            chosen = exact[0]
        elif len(exact) > 1:
            # 유형 문자열 겹침 우선, 그 다음 면적 최대
            biz_type = row.get("BIZ_TYPE") or ""
            scored = sorted(
                exact,
                key=lambda f: (
                    0 if (f["properties"].get("type") and f["properties"]["type"] in biz_type) else 1,
                    -(f["properties"].get("area") or 0),
                ),
            )
            chosen = scored[0]
            if len(exact) >= 2 and len(ambiguous_samples) < 20:
                ambiguous_samples.append(
                    {
                        "stat": {"seq_no": seq_no, "gu": gu, "zone_nm": zone_nm},
                        "candidates": [f["properties"]["name"] for f in exact],
                        "chosen": chosen["properties"]["name"],
                        "kind": "exact-multi",
                    }
                )
        elif len(contain) == 1:
            chosen = contain[0]
        elif len(contain) > 1:
            biz_type = row.get("BIZ_TYPE") or ""
            scored = sorted(
                contain,
                key=lambda f: (
                    0 if (f["properties"].get("type") and f["properties"]["type"] in biz_type) else 1,
                    -(f["properties"].get("area") or 0),
                ),
            )
            chosen = scored[0]
            if len(ambiguous_samples) < 20:
                ambiguous_samples.append(
                    {
                        "stat": {"seq_no": seq_no, "gu": gu, "zone_nm": zone_nm},
                        "candidates": [f["properties"]["name"] for f in contain],
                        "chosen": chosen["properties"]["name"],
                        "kind": "contain-multi",
                    }
                )

        if chosen:
            matched += 1
            zone_id = chosen["properties"]["id"]
            stat_zone_id[seq_no] = zone_id
            chosen["properties"]["statId"] = seq_no
        else:
            stat_zone_id[seq_no] = None
            if len(unmatched_samples) < 20:
                unmatched_samples.append({"seq_no": seq_no, "gu": gu, "zone_nm": zone_nm, "jibun": row.get("JIBUN_ADDR")})

    total = len(stat_rows)
    report = {
        "total": total,
        "matched": matched,
        "match_rate": round(matched / total, 4) if total else 0,
        "unmatched_samples": unmatched_samples,
        "ambiguous_samples": ambiguous_samples,
    }
    return stat_zone_id, report


# ---------------------------------------------------------------------------
# 조인 2: 입찰공고/협력업체(BIZ_NO -> ASCT_NM) <-> 구역
# ---------------------------------------------------------------------------


def join_biz_to_zones(biz_to_asct: dict[str, str], zone_features):
    all_zones_by_norm: dict[str, list[dict]] = defaultdict(list)
    for feat in zone_features:
        norm = normalize_zone_name(feat["properties"]["name"])
        if norm:
            all_zones_by_norm[norm].append(feat)

    biz_zone_id: dict[str, str | None] = {}
    matched = 0
    for biz_no, asct_nm in biz_to_asct.items():
        norm_target = normalize_zone_name(asct_nm)
        chosen = None
        if norm_target:
            exact = all_zones_by_norm.get(norm_target, [])
            if len(exact) == 1:
                chosen = exact[0]
            elif len(exact) == 0:
                contain = []
                for norm_cand, feats in all_zones_by_norm.items():
                    if norm_target in norm_cand or norm_cand in norm_target:
                        contain.extend(feats)
                if len(contain) == 1:
                    chosen = contain[0]
                # 다중 후보면 매칭 포기 (len(contain) > 1 -> None)
            # exact 다중 후보도 자치구 정보 없으므로 포기

        if chosen:
            matched += 1
            zone_id = chosen["properties"]["id"]
            biz_zone_id[biz_no] = zone_id
            if biz_no not in chosen["properties"]["bizNos"]:
                chosen["properties"]["bizNos"].append(biz_no)
        else:
            biz_zone_id[biz_no] = None

    total = len(biz_to_asct)
    report = {
        "total": total,
        "matched": matched,
        "match_rate": round(matched / total, 4) if total else 0,
    }
    return biz_zone_id, report


# ---------------------------------------------------------------------------
# 출력 빌더
# ---------------------------------------------------------------------------


def build_projects(stat_rows, stat_zone_id):
    projects = []
    for row in stat_rows:
        seq_no = row.get("SEQ_NO")

        def g(key):
            v = row.get(key)
            return v if v not in (None, "") else None

        households = {}
        if g("EXISTING_HOUSEHOLDS") is not None:
            households["existing"] = g("EXISTING_HOUSEHOLDS")
        if g("TOT_BUILT_HOUSEHOLDS") is not None:
            households["total"] = g("TOT_BUILT_HOUSEHOLDS")
        if g("SALE_BUILT_HOUSEHOLDS") is not None:
            households["sale"] = g("SALE_BUILT_HOUSEHOLDS")
        if g("RENT_BUILT_HOUSEHOLDS") is not None:
            households["rent"] = g("RENT_BUILT_HOUSEHOLDS")

        dates = {}
        date_map = {
            "zoneDesignation": ("ZONE_DESIGNATION_INIT_YMD", "ZONE_DESIGNATION_LAST_YMD"),
            "committee": ("PROMOTION_COMMITTEE_YMD", None),
            "association": ("ASSOCIATION_ESTABLISHMENT_YMD", None),
            "archReview": ("ARCHITECTURAL_REVIEW_YMD", None),
            "bizImpl": ("BIZ_IMPLEMENTATION_INIT_YMD", "BIZ_IMPLEMENTATION_LAST_YMD"),
            "mgmtDisposition": ("MGMT_DISPOSITION_INIT_YMD", "MGMT_DISPOSITION_LAST_YMD"),
            "migrationStart": ("MIGRATION_START_YMD", None),
            "migrationEnd": ("MIGRATION_END_YMD", None),
            "construction": ("CONSTRUCTION_START_YMD", None),
        }
        for out_key, (init_key, last_key) in date_map.items():
            val = g(init_key)
            if val is None and last_key:
                val = g(last_key)
            if val is not None:
                dates[out_key] = val

        project = {
            "id": seq_no,
            "zoneId": stat_zone_id.get(seq_no),
            "name": row.get("ZONE_NM"),
            "gu": row.get("DISTRICT"),
            "jibun": g("JIBUN_ADDR"),
            "road": g("ROAD_ADDR"),
            "publicPrivate": g("BIZ_METHOD_PUBLIC_PRIVATE"),
            "generalPromoted": g("BIZ_METHOD_GENERAL_PROMOTED"),
            "type": g("BIZ_TYPE"),
            "stage": g("BIZ_STAGE"),
        }
        if households:
            project["households"] = households
        if dates:
            project["dates"] = dates
        projects.append(project)
    return projects


def parse_ymd_any(s: str):
    s = (s or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def build_bids(bid_rows):
    by_biz: dict[str, list[dict]] = defaultdict(list)
    asct_by_biz: dict[str, str] = {}
    for row in bid_rows:
        biz_no = row.get("BIZ_NO")
        if not biz_no:
            continue
        asct_nm = row.get("ASCT_NM") or ""
        if asct_nm and biz_no not in asct_by_biz:
            asct_by_biz[biz_no] = asct_nm
        entry = {
            "ttl": row.get("TTL") or None,
            "asct": asct_nm or None,
            "pbanc": row.get("PBANC_YMD") or None,
            "bidNm": row.get("BID_NM") or None,
            "expiry": row.get("BID_EXPRY_DAY") or None,
            "method": row.get("SLCTN_MTHD") or None,
            "_sort": row.get("REG_DAY") or row.get("PBANC_YMD") or "",
        }
        by_biz[biz_no].append(entry)

    out = {}
    for biz_no, entries in by_biz.items():
        entries.sort(key=lambda e: e["_sort"], reverse=True)
        for e in entries:
            e.pop("_sort", None)
        out[biz_no] = entries[:30]
    return out, asct_by_biz


def build_partners(partner_rows):
    out: dict[str, list[dict]] = defaultdict(list)
    for row in partner_rows:
        biz_no = row.get("BIZ_NO")
        if not biz_no:
            continue
        out[biz_no].append(
            {
                "cls": row.get("BZENTY_CLSF_NM") or None,
                "name": row.get("BZENTY_NM") or None,
                "cancelled": (row.get("CTRT_CNCLTN_YN") or "N").strip().upper() == "Y",
            }
        )
    return out


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def main():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    api_key = load_seoul_api_key()

    print("== 1) 구역 폴리곤 + 코드표 파싱 ==")
    members = extract_zip_members(ZIP_PATH)
    xlsx_bytes = members["서울플랜+코드정의표(대민용).xlsx"]
    lclas_label, sclas_label, propel_label, all_propel_codes = parse_code_tables(xlsx_bytes)
    prj_text = members["UPIS_C_UQ120.prj"].decode("utf-8")

    zone_features, shp_stats = build_zones(
        members["UPIS_C_UQ120.shp"],
        members["UPIS_C_UQ120.shx"],
        members["UPIS_C_UQ120.dbf"],
        prj_text,
        lclas_label,
        sclas_label,
        propel_label,
    )
    print(f"  구역 폴리곤 {len(zone_features)}건 파싱 완료")
    if shp_stats["missing_propel_codes"]:
        print(f"  [경고] 코드표에 없는 PROPEL_CD: {shp_stats['missing_propel_codes']}")

    print("== 2) 정비사업 통계 (TbSeoulRedevStatus) 수집 ==")
    stat_rows = fetch_all(api_key, "TbSeoulRedevStatus", total_hint=496)

    print("== 3) 조합 입찰공고 (ListCombinationBidding) 수집 ==")
    bid_rows = fetch_all(api_key, "ListCombinationBidding", total_hint=14763)

    print("== 4) 협력업체 (CleanupPartnersInfo) 수집 ==")
    partner_rows = fetch_all(api_key, "CleanupPartnersInfo")

    print("== 5) 조인 1: 통계 <-> 폴리곤 ==")
    stat_zone_id, join1_report = join_stats_to_zones(stat_rows, zone_features)
    print(f"  매칭률: {join1_report['matched']}/{join1_report['total']} = {join1_report['match_rate']:.2%}")

    print("== 6) 조인 2: 입찰/협력업체 <-> 구역 (BIZ_NO -> ASCT_NM) ==")
    bids_by_biz, asct_by_biz = build_bids(bid_rows)
    biz_zone_id, join2_report = join_biz_to_zones(asct_by_biz, zone_features)
    print(f"  매칭률: {join2_report['matched']}/{join2_report['total']} = {join2_report['match_rate']:.2%}")

    partners_by_biz = build_partners(partner_rows)

    print("== 7) 출력 작성 ==")
    zones_geojson = {"type": "FeatureCollection", "features": zone_features}
    projects = build_projects(stat_rows, stat_zone_id)
    bizzones = {biz_no: {"zoneId": biz_zone_id.get(biz_no), "asctNm": asct_by_biz.get(biz_no)} for biz_no in asct_by_biz}

    write_json(DATA_DIR / "zones.geojson", zones_geojson)
    write_json(DATA_DIR / "projects.json", projects)
    write_json(DATA_DIR / "bids.json", bids_by_biz)
    write_json(DATA_DIR / "bizzones.json", bizzones)
    write_json(DATA_DIR / "partners.json", partners_by_biz)

    meta = {
        "collectedAt": datetime.now(timezone.utc).isoformat(),
        "sources": {
            "zones": {"file": str(ZIP_PATH.relative_to(ROOT)), "count": len(zone_features)},
            "projects": {"service": "TbSeoulRedevStatus", "count": len(stat_rows)},
            "bids": {"service": "ListCombinationBidding", "count": len(bid_rows), "distinctBizNo": len(bids_by_biz)},
            "partners": {"service": "CleanupPartnersInfo", "count": len(partner_rows), "distinctBizNo": len(partners_by_biz)},
        },
        "joins": {
            "statsToZones": {k: v for k, v in join1_report.items() if k not in ("unmatched_samples", "ambiguous_samples")},
            "bizToZones": join2_report,
        },
        "codeTable": {
            "missingPropelCodes": shp_stats["missing_propel_codes"],
        },
        "simplify": {
            "beforeBytes": shp_stats["simplify_before_bytes"],
            "afterBytes": shp_stats["simplify_after_bytes"],
            "topologyBreaks": shp_stats["topology_breaks"],
        },
    }
    write_json(DATA_DIR / "meta.json", meta)

    # 리포트 출력용 샘플 저장 (data/ 에는 포함 안 함)
    report_path = ROOT / "scripts" / "raw" / "join_report.json"
    write_json(report_path, {"join1": join1_report, "join2": join2_report})

    print_final_report(zones_geojson, projects, bids_by_biz, partners_by_biz, meta, join1_report, join2_report)


def write_json(path: Path, obj) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def print_final_report(zones_geojson, projects, bids_by_biz, partners_by_biz, meta, join1_report, join2_report):
    import gzip

    def size_of(path: Path) -> int:
        return path.stat().st_size

    def gzip_size(path: Path) -> int:
        return len(gzip.compress(path.read_bytes(), compresslevel=9))

    print("\n================ 최종 보고 ================")
    for name in ["zones.geojson", "projects.json", "bids.json", "bizzones.json", "partners.json", "meta.json"]:
        p = DATA_DIR / name
        if p.exists():
            print(f"{name}: {size_of(p):,} bytes (gzip {gzip_size(p):,} bytes)")

    print(f"\nzones: {len(zones_geojson['features'])}건")
    print(f"projects: {len(projects)}건")
    print(f"bids(biz_no 단위): {len(bids_by_biz)}건")
    print(f"partners(biz_no 단위): {len(partners_by_biz)}건")

    print(f"\n단순화 전/후: {meta['simplify']['beforeBytes']:,} -> {meta['simplify']['afterBytes']:,} bytes "
          f"(위상 깨짐으로 원본 유지: {meta['simplify']['topologyBreaks']}건)")

    print(f"\n[조인1] 통계<->폴리곤 매칭률: {join1_report['matched']}/{join1_report['total']} = {join1_report['match_rate']:.2%}")
    print("  미매칭 샘플(최대 20):")
    for s in join1_report["unmatched_samples"]:
        print(f"    - {s}")
    print("  오매칭 의심 샘플(최대 20):")
    for s in join1_report["ambiguous_samples"]:
        print(f"    - {s}")

    print(f"\n[조인2] 입찰/협력업체<->구역 매칭률: {join2_report['matched']}/{join2_report['total']} = {join2_report['match_rate']:.2%}")

    if meta["codeTable"]["missingPropelCodes"]:
        print(f"\n코드표 누락 PROPEL_CD: {meta['codeTable']['missingPropelCodes']}")
    else:
        print("\n코드표 누락 PROPEL_CD: 없음 (전량 매칭)")


if __name__ == "__main__":
    main()
